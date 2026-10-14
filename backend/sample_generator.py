"""
Generates realistic synthetic Federal FOIA sample documents (PDF, DOCX, TXT)
for testing and demonstration of the Offline FOIA Redaction Workbench.
All names, phone numbers, email addresses, SSNs, and badge numbers are 100% fictional.
"""

import io
from pathlib import Path
import fitz
import docx
from docx.shared import Pt, Inches
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas as rl_canvas


def generate_all_sample_documents(sample_dir: Path) -> None:
    sample_dir.mkdir(parents=True, exist_ok=True)
    _create_sample_pdf(sample_dir / "DHS_OIG_Investigative_Memorandum_2026-0418.pdf")
    _create_sample_scanned_pdf_and_png(
        sample_dir / "VA_OIG_Scanned_Field_Intake_2026_IMAGE_SCAN.pdf",
        sample_dir / "VA_OIG_Scanned_Field_Intake_2026_SCAN.png",
    )
    _create_sample_docx(sample_dir / "DOJ_FOIA_Witness_Interview_Summary_2026.docx")
    _create_sample_txt(sample_dir / "GSA_Personnel_Incident_Correspondence_Log.txt")


def _create_sample_scanned_pdf_and_png(scanned_pdf_path: Path, scanned_png_path: Path) -> None:
    """
    Creates a realistic image-only scanned PDF (0 selectable text characters, pure raster image)
    and a standalone scanned PNG image to demonstrate offline Tesseract OCR + pixel-level redaction.
    """
    buf = io.BytesIO()
    c = rl_canvas.Canvas(buf, pagesize=letter)
    width, height = letter
    margin = 54

    y = height - 55
    c.setFont("Helvetica-Bold", 11)
    c.drawCentredString(width / 2, y, "UNCLASSIFIED // SCANNED FIELD RECORD (FLATBED DIGITIZED COPY)")
    y -= 26

    c.setFont("Helvetica-Bold", 14)
    c.drawString(margin, y, "DEPARTMENT OF VETERANS AFFAIRS — OFFICE OF INSPECTOR GENERAL")
    y -= 18
    c.setFont("Helvetica-Bold", 11)
    c.drawString(margin, y, "SCANNED HOTLINE INTAKE & WITNESS RECORD — FILE #VA-IG-2026-3310")
    y -= 12
    c.line(margin, y, width - margin, y)
    y -= 24

    c.setFont("Helvetica", 10.5)
    scan_lines = [
        "DATE RECEIVED:    October 14, 2026",
        "INTAKE OFFICER:   Special Agent Valerie M. Sterling",
        "AGENT CONTACT:    Tel: (202) 555-0149   |   Email: valerie.sterling@va.gov",
        "COMPLAINANT:      Dr. Thomas H. Abernathy, Chief Medical Officer",
        "DIRECT PHONE:     (202) 555-0198 ext. 302   |   Personal Cell: 703-555-0166",
        "OFFICIAL EMAIL:   thomas.abernathy@va.gov   |   Personal: t.abernathy.md@gmail.com",
        "VERIFICATION SSN: 441-29-8831   |   CREDENTIAL: Badge #VA-902145",
        "HOME ADDRESS:     1640 Riverside Drive, Apt 3A, Arlington, VA 22201",
        "",
        "SUMMARY OF SCANNED FIELD REPORT:",
        "Dr. Thomas H. Abernathy reported that pharmacy logistics supervisor",
        "Mr. Gregory P. Donovan (Phone: (804) 555-0174, Email: gregory.donovan@va.gov)",
        "and contractor representative Ms. Nadia K. Al-Mansoor (Cell: (571) 555-0181,",
        "Email: nadia.almansoor@medsupply-fed.com, SSN: 519-40-6623) were listed on",
        "unredacted shipping manifests at the Richmond Regional Medical Center.",
        "",
        "FOIA ANALYST NOTE: This document was scanned from a paper hardcopy with no",
        "embedded digital text. Offline Tesseract OCR is required to detect and excise",
        "both the OCR text layer and underlying raster pixels under Exemption (b)(6).",
    ]
    for line in scan_lines:
        c.drawString(margin, y, line)
        y -= 16

    c.save()

    # Rasterize the vector PDF into a 200 DPI image and save both a PNG and an image-only PDF
    temp_vector_doc = fitz.open("pdf", buf.getvalue())
    pix = temp_vector_doc[0].get_pixmap(dpi=200, alpha=False)
    png_bytes = pix.tobytes("png")
    temp_vector_doc.close()

    scanned_png_path.write_bytes(png_bytes)

    img_only_pdf = fitz.open()
    page = img_only_pdf.new_page(width=width, height=height)
    page.insert_image(page.rect, stream=png_bytes)
    img_only_pdf.save(str(scanned_pdf_path), deflate=True)
    img_only_pdf.close()


def _create_sample_pdf(path: Path) -> None:
    c = rl_canvas.Canvas(str(path), pagesize=letter)
    width, height = letter
    margin = 54

    # ---------------- PAGE 1 ----------------
    y = height - 50
    c.setFont("Helvetica-Bold", 11)
    c.drawCentredString(width / 2, y, "UNCLASSIFIED // FOR OFFICIAL USE ONLY (FOUO)")
    y -= 24

    c.setFont("Helvetica-Bold", 14)
    c.drawString(margin, y, "OFFICE OF INSPECTOR GENERAL — DEPARTMENT OF HOMELAND SECURITY")
    y -= 18
    c.setFont("Helvetica-Bold", 11)
    c.drawString(margin, y, "INVESTIGATIVE MEMORANDUM — CASE FILE #OIG-INV-2026-0418")
    y -= 10
    c.line(margin, y, width - margin, y)
    y -= 20

    c.setFont("Helvetica", 10)
    memo_headers = [
        "DATE:          October 14, 2026",
        "MEMORANDUM FOR: Evelyn M. Vance, Deputy Assistant Inspector General for Investigations",
        "FROM:          Special Agent Marcus T. Holloway, Field Office VII (Arlington)",
        "SUBJECT:       Alleged Improper Contractor Access & Privacy Incident — Facility Sector 4",
        "POINT OF CONTACT: Sarah K. Jenkins | Tel: (202) 555-0184 ext. 442 | Email: sarah.jenkins@oig.dhs.gov",
    ]
    for line in memo_headers:
        c.drawString(margin, y, line)
        y -= 15

    y -= 6
    c.line(margin, y, width - margin, y)
    y -= 20

    c.setFont("Helvetica-Bold", 11)
    c.drawString(margin, y, "1. EXECUTIVE SUMMARY & COMPLAINANT DETAILS")
    y -= 16

    c.setFont("Helvetica", 10)
    p1_lines = [
        "On September 29, 2026, the OIG Hotline received a formal disclosure from Complainant",
        "Dr. Elena V. Rostova, Senior Systems Architect at the National Infrastructure Operations",
        "Center. Dr. Elena V. Rostova reported unauthorized after-hours badge access in Server Vault B.",
        "Complainant's direct government phone is (571) 555-0199, personal mobile is 703-555-0147,",
        "and personal email address is elena.rostova.phd@protonmail.com (work: elena.rostova@hq.dhs.gov).",
        "Her residential address on file is 4821 Old Dominion Drive, Apt 4B, Arlington, VA 22207,",
        "and her employee verification SSN on the intake form is 412-58-9034 (EDIPI: 1094827361).",
    ]
    for line in p1_lines:
        c.drawString(margin, y, line)
        y -= 14

    y -= 12
    c.setFont("Helvetica-Bold", 11)
    c.drawString(margin, y, "2. SUBJECTS & WITNESS CONTACTS INTERVIEWED")
    y -= 16

    c.setFont("Helvetica", 10)
    p2_lines = [
        "Field investigators Special Agent Marcus T. Holloway (Phone: +1 (202) 555-0162, Email:",
        "marcus.holloway@oig.dhs.gov) and Auditor David R. Chen (Tel: 202.555.0128, Email:",
        "david.chen2@oig.dhs.gov) conducted sworn interviews with the following personnel:",
        "",
        "  a) Witness #1: Capt. Jonathan P. Mercer, Facility Security Officer",
        "     - Military DSN: DSN 312-555-8821  |  Commercial: (703) 555-0173",
        "     - Official Email: jonathan.p.mercer.mil@mail.mil",
        "     - Home Address: 1190 Braddock Road, Alexandria, VA 22302",
        "",
        "  b) Witness #2: Ms. Priya N. Patel, Lead Network Engineer (Contractor)",
        "     - Cell Phone: (301) 555-0156  |  Office Tel: (202) 555-0191 x308",
        "     - Email Address: priya.patel@apex-fed-solutions.com",
        "     - Badge Number: Badge #CTR-884920  |  SSN: 503-44-7812",
        "",
        "  c) Subject Individual: Mr. Viktor S. Lindqvist, Senior Database Administrator",
        "     - Direct Phone: (410) 555-0133  |  Alternate Cell: 240-555-0188",
        "     - Personal Email: v.lindqvist77@gmail.com  |  Gov Email: viktor.lindqvist@cisa.dhs.gov",
        "     - Home Address: 742 Evergreen Terrace, Suite 200, Bethesda, MD 20814",
    ]
    for line in p2_lines:
        c.drawString(margin, y, line)
        y -= 14

    y -= 12
    c.setFont("Helvetica-Bold", 11)
    c.drawString(margin, y, "3. FINDINGS & PRIVACY ACT / FOIA RELEASE GUIDANCE")
    y -= 16
    c.setFont("Helvetica", 10)
    p3_lines = [
        "Access logs confirm that Mr. Viktor S. Lindqvist entered Vault B at 23:14 EST under an",
        "emergency maintenance ticket approved verbally by Branch Chief Robert L. Alvarez",
        "(robert.alvarez@cisa.dhs.gov, Tel: (202) 555-0115). No classified data exfiltration occurred.",
        "Prior to any public FOIA release under 5 U.S.C. § 552, all personal names, phone numbers,",
        "email addresses, SSNs, and home addresses must be redacted under Exemptions (b)(6) and (b)(7)(C).",
    ]
    for line in p3_lines:
        c.drawString(margin, y, line)
        y -= 14

    c.setFont("Helvetica-Bold", 9)
    c.drawCentredString(width / 2, 35, "PAGE 1 OF 2 — UNCLASSIFIED // FOUO — SUBJECT TO FOIA EXEMPTIONS (b)(6) & (b)(7)(C)")
    c.showPage()

    # ---------------- PAGE 2 ----------------
    y = height - 50
    c.setFont("Helvetica-Bold", 11)
    c.drawCentredString(width / 2, y, "UNCLASSIFIED // FOR OFFICIAL USE ONLY (FOUO)")
    y -= 28

    c.setFont("Helvetica-Bold", 12)
    c.drawString(margin, y, "APPENDIX A: CALL LOG & EMAIL CORRESPONDENCE EXHIBIT (#OIG-INV-2026-0418)")
    y -= 14
    c.line(margin, y, width - margin, y)
    y -= 20

    c.setFont("Helvetica", 10)
    page2_lines = [
        "Exhibit A-1: Extracted Emergency Notification Chain",
        "--------------------------------------------------------------------------------",
        "From: Viktor S. Lindqvist <viktor.lindqvist@cisa.dhs.gov>",
        "To: Robert L. Alvarez <robert.alvarez@cisa.dhs.gov>; Priya N. Patel <priya.patel@apex-fed-solutions.com>",
        "Cc: Elena V. Rostova <elena.rostova@hq.dhs.gov>; Evelyn M. Vance <evelyn.vance@oig.dhs.gov>",
        "Subject: URGENT: Cluster Failover in Vault B",
        "",
        "Robert and Priya,",
        "I am heading into Vault B now to patch the primary node. If the console drops, please",
        "call my personal cell immediately at (410) 555-0133 or my backup line at 240-555-0188.",
        "On-call technician Sgt. Carlos M. Mendoza can also be reached at (703) 555-0164 or",
        "via email at carlos.mendoza@uscg.mil.",
        "",
        "Signed,",
        "Special Agent Marcus T. Holloway",
        "Office of Inspector General | Tel: (202) 555-0162 | marcus.holloway@oig.dhs.gov",
    ]
    for line in page2_lines:
        c.drawString(margin, y, line)
        y -= 15

    c.setFont("Helvetica-Bold", 9)
    c.drawCentredString(width / 2, 35, "PAGE 2 OF 2 — UNCLASSIFIED // FOUO — SUBJECT TO FOIA EXEMPTIONS (b)(6) & (b)(7)(C)")
    c.save()


def _create_sample_docx(path: Path) -> None:
    d = docx.Document()
    d.core_properties.author = "Special Agent Hannah R. Abbott"
    d.core_properties.title = "DOJ FOIA Witness Interview Summary"

    h = d.add_heading("U.S. DEPARTMENT OF JUSTICE — FOIA REVIEW RECORD", level=1)
    p_sub = d.add_paragraph("MEMORANDUM OF WITNESS INTERVIEW — FILE #DOJ-CIV-2026-8821")
    p_sub.runs[0].bold = True

    d.add_paragraph(
        "DATE: October 10, 2026\n"
        "INTERVIEWING AGENT: Special Agent Hannah R. Abbott (Tel: (202) 555-0141, Email: hannah.abbott@usdoj.gov)\n"
        "WITNESS NAME: Mr. Nathaniel J. Brooks, Regional Procurement Officer\n"
        "WITNESS CONTACT: Work (312) 555-0192 ext. 114 | Cell 312-555-0177 | Email: nathaniel.brooks@gsa.gov\n"
        "PERSONAL EMAIL: nate.brooks1981@outlook.com | SSN: 389-62-1145\n"
        "RESIDENTIAL ADDRESS: 2940 North Lake Shore Drive, Apt 12C, Chicago, IL 60657"
    )

    d.add_heading("1. Summary of Testimony", level=2)
    d.add_paragraph(
        "On October 8, 2026, Special Agent Hannah R. Abbott and Trial Attorney Benjamin C. Sterling "
        "(benjamin.sterling@usdoj.gov, (202) 555-0155) interviewed Mr. Nathaniel J. Brooks regarding "
        "Contract Modification #GS-05P-26-D-0092. Mr. Nathaniel J. Brooks stated that vendor account "
        "representative Ms. Olivia M. Hawthorne (Email: olivia.hawthorne@vanguard-logistics.org, "
        "Phone: (404) 555-0123) contacted him on his personal telephone at 312-555-0177."
    )

    d.add_heading("2. Personnel & Contact Table", level=2)
    table = d.add_table(rows=1, cols=4)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Full Name"
    hdr_cells[1].text = "Role / Title"
    hdr_cells[2].text = "Telephone Number"
    hdr_cells[3].text = "Email Address"

    rows_data = [
        ("Nathaniel J. Brooks", "Procurement Officer", "(312) 555-0192 ext. 114", "nathaniel.brooks@gsa.gov"),
        ("Olivia M. Hawthorne", "Vendor Account Exec", "(404) 555-0123", "olivia.hawthorne@vanguard-logistics.org"),
        ("Dr. Raymond K. Mercer", "Technical Evaluator", "(202) 555-0189", "raymond.mercer@nist.gov"),
        ("Hannah R. Abbott", "Special Agent", "(202) 555-0141", "hannah.abbott@usdoj.gov"),
    ]
    for name, role, phone, email in rows_data:
        row_cells = table.add_row().cells
        row_cells[0].text = name
        row_cells[1].text = role
        row_cells[2].text = phone
        row_cells[3].text = email

    d.save(str(path))


def _create_sample_txt(path: Path) -> None:
    content = """================================================================================
FEDERAL GENERAL SERVICES ADMINISTRATION (GSA) — FOIA RELEASE CANDIDATE
DOCUMENT ID: GSA-FOIA-2026-09912  |  CLASSIFICATION: UNCLASSIFIED // CUI
================================================================================

FROM:      Director Margaret A. Sullivan
EMAIL:     margaret.sullivan@gsa.gov
PHONE:     (202) 555-0104 / Mobile: (571) 555-0122
TO:        FOIA Public Liaison Officer Arthur D. Pendelton (arthur.pendelton@gsa.gov)
DATE:      October 12, 2026
SUBJECT:   Citizen Correspondence & Personnel Roster Review for FOIA Request #2026-00419

--------------------------------------------------------------------------------
1. CITIZEN REQUESTER & THIRD-PARTY PRIVACY SUMMARY
--------------------------------------------------------------------------------
The underlying administrative record contains third-party personally identifiable
information (PII) that must be withheld pursuant to 5 U.S.C. § 552(b)(6) prior to
public disclosure:

  * Field Inspector: Mr. Gabriel L. Morales
    - Direct Office Phone: (415) 555-0168 ext. 204
    - Personal Cell Phone: 650-555-0194
    - Official Email: gabriel.morales@gsa.gov
    - Personal Email: gmorales.sf@gmail.com
    - Social Security Number: 548-31-8820
    - Home Address: 845 Market Street, Apt 902, San Francisco, CA 94103

  * External Whistleblower / Witness: Ms. Clara E. Lindstrom
    - Telephone: +1-206-555-0139
    - Email Address: clara.lindstrom@northwest-audits.com
    - Badge ID: Badge #GSA-773190

  * Reviewing Supervisor: Col. Harrison W. Vance
    - Military DSN: DSN 421-8890  |  Commercial: (703) 555-0118
    - Email: harrison.w.vance.mil@mail.mil

--------------------------------------------------------------------------------
2. DETERMINATION
--------------------------------------------------------------------------------
Please ensure all names, phone numbers, email addresses, SSNs, badge IDs, and
residential addresses above are completely excised from the final release file
and logged in the accompanying Vaughn Index under FOIA Exemption (b)(6).
"""
    path.write_text(content, encoding="utf-8")
